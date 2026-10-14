"""python-docx helpers: tables with dual widths, captions with SEQ fields,
a TOC field, page-number footers, landscape sections."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEADER_FILL = "D9E2F3"     # the light blue the assignment brief uses for table headers
ALT_FILL = "F5F7FB"
INK = RGBColor(0x1F, 0x2A, 0x44)

STATUS_FILL = {"Pass": "E2F0D9", "Fail": "FBE3D6", "Blocked": "FFF2CC", "Not Executed": "EDEDED"}


class Doc:
    def __init__(self):
        self.d = Document()
        self.tables = 0
        self.figures = 0
        self.headings = []
        self._styles()

    # ---------------------------------------------------------------- setup
    def _styles(self):
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for name, size, before in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11.5, 10)):
            h = st[name]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = INK
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.bold = False
        cap.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.0))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        # Ask Word / LibreOffice to refresh fields (TOC, SEQ, NUMPAGES) on open.
        settings = self.d.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    # ------------------------------------------------------------ primitives
    def h(self, text, level=1):
        if level <= 2:
            self.headings.append((level, text))
        return self.d.add_heading(text, level=level)

    def p(self, text="", bold=False, italic=False, size=None, align=None, after=None, color=None, style=None):
        para = self.d.add_paragraph(style=style)
        if text:
            self._runs(para, text, bold=bold, italic=italic, size=size, color=color)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    @staticmethod
    def _runs(para, text, bold=False, italic=False, size=None, color=None):
        """Supports **bold** spans inside text."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = para.add_run(part)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return para

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            self._runs(para, it)
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --------------------------------------------------------------- fields
    @staticmethod
    def _field(para, instr, cached="", bold=False, size=None):
        def run_with(child):
            r = OxmlElement("w:r")
            if bold or size:
                rpr = OxmlElement("w:rPr")
                if bold:
                    rpr.append(OxmlElement("w:b"))
                if size:
                    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
                r.append(rpr)
            r.append(child)
            para._p.append(r)
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); run_with(b)
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run_with(it)
        s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); run_with(s)
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = cached; run_with(t)
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run_with(e)

    def toc(self, entries=None):
        """entries: [(level, text, page)] written as the field's cached result, so the
        contents read correctly before any field update; Word can still refresh it."""
        if not entries:
            para = self.d.add_paragraph()
            self._field(para, r'TOC \o "1-2" \h \z \u',
                        "Right-click here and choose Update Field to build the table of contents.")
            return

        def fld(para, kind):
            r = OxmlElement("w:r"); c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), kind)
            r.append(c); para._p.append(r)

        for i, (level, text, page) in enumerate(entries):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(2 if level == 2 else 4)
            pf.space_before = Pt(0 if level == 2 else 4)
            pf.left_indent = Cm(0 if level == 1 else 0.6)
            pf.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if i == 0:
                fld(para, "begin")
                r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
                it.text = r' TOC \o "1-2" \h \z \u '; r.append(it); para._p.append(r)
                fld(para, "separate")
            run = para.add_run(f"{text}\t{page}")
            run.font.size = Pt(10.5 if level == 1 else 10)
            run.bold = level == 1
            if i == len(entries) - 1:
                fld(para, "end")

    def caption(self, kind, text):
        """kind = 'Table' or 'Figure'. Numbered with a SEQ field whose cached value is correct."""
        if kind == "Table":
            self.tables += 1; n = self.tables
        else:
            self.figures += 1; n = self.figures
        para = self.d.add_paragraph(style="Caption")
        para.add_run(f"{kind} ")
        self._field(para, f"SEQ {kind} \\* ARABIC", str(n))
        para.add_run(f": {text}")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "Figure" else WD_ALIGN_PARAGRAPH.LEFT
        if kind == "Table":
            # Called after the table was added: move the caption to sit above it.
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(6)
            node = para._p.getprevious()
            while node is not None and node.tag != qn("w:tbl"):
                node = node.getprevious()
            if node is not None:
                node.addprevious(para._p)
        return n

    def footer(self, section, left_text):
        f = section.footer
        para = f.paragraphs[0]
        para.text = ""
        r = para.add_run(left_text + "    |    Page "); r.font.size = Pt(8.5)
        self._field(para, "PAGE", "1", size=8.5)
        r = para.add_run(" of "); r.font.size = Pt(8.5)
        self._field(para, "NUMPAGES", "1", size=8.5)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def header(self, section, text):
        para = section.header.paragraphs[0]
        para.text = ""
        r = para.add_run(text); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # -------------------------------------------------------------- sections
    def new_section(self, landscape=False):
        sec = self.d.add_section()
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0 if not landscape else 1.6)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        return sec

    # ---------------------------------------------------------------- tables
    @staticmethod
    def _shade(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    @staticmethod
    def _cell_width(cell, cm):
        tcpr = cell._tc.get_or_add_tcPr()
        w = OxmlElement("w:tcW"); w.set(qn("w:w"), str(int(cm * 567))); w.set(qn("w:type"), "dxa")
        tcpr.append(w)

    @staticmethod
    def _repeat_header(row):
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trpr.append(el)

    @staticmethod
    def _cant_split(row):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))

    def table(self, header, rows, widths_cm, size=9, zebra=True, status_col=None, bold_first_col=False,
              header_fill=HEADER_FILL, align_center_cols=()):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        tblpr = t._tbl.tblPr
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths_cm) * 567))); tw.set(qn("w:type"), "dxa")
        tblpr.append(tw)
        hdr = t.rows[0]
        self._repeat_header(hdr)
        for i, (cell, text) in enumerate(zip(hdr.cells, header)):
            self._cell_width(cell, widths_cm[i]); self._shade(cell, header_fill)
            para = cell.paragraphs[0]; para.paragraph_format.space_after = Pt(0)
            r = para.add_run(str(text)); r.bold = True; r.font.size = Pt(size)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            self._cant_split(t.rows[-1])
            for i, (cell, val) in enumerate(zip(cells, row)):
                self._cell_width(cell, widths_cm[i])
                lines = val if isinstance(val, list) else [val]
                for li, line in enumerate(lines):
                    para = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.0
                    self._runs(para, "" if line is None else str(line), size=size,
                               bold=(bold_first_col and i == 0))
                    if i in align_center_cols:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if status_col is not None and i == status_col and str(val) in STATUS_FILL:
                    self._shade(cell, STATUS_FILL[str(val)])
                elif zebra and ri % 2 == 1:
                    self._shade(cell, ALT_FILL)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, pairs, widths_cm=(4.2, 12.8), size=9.5, status_key="Status"):
        """Two-column field/value block (used for each test case and defect report)."""
        t = self.d.add_table(rows=0, cols=2)
        t.style = "Table Grid"; t.autofit = False
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
        for key, val in pairs:
            row = t.add_row(); self._cant_split(row)
            k, v = row.cells
            self._cell_width(k, widths_cm[0]); self._cell_width(v, widths_cm[1]); self._shade(k, HEADER_FILL)
            kp = k.paragraphs[0]; kp.paragraph_format.space_after = Pt(0)
            r = kp.add_run(key); r.bold = True; r.font.size = Pt(size)
            lines = val if isinstance(val, list) else [val]
            for li, line in enumerate(lines):
                para = v.paragraphs[0] if li == 0 else v.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                self._runs(para, str(line), size=size)
            if key == status_key and str(val) in STATUS_FILL:
                self._shade(v, STATUS_FILL[str(val)])
        for row in t.rows[:-1]:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def kv_table_image(self, pairs, image, image_cm, widths_cm=(3.0, 9.4), size=8.5, caption=None):
        """Field/value block with one screenshot beside it in a merged third column (compact defect report)."""
        t = self.kv_table(pairs, widths_cm=widths_cm, size=size, status_key=None)
        spacer = t._tbl.getnext(); spacer.getparent().remove(spacer)
        grid = t._tbl.tblGrid
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(image_cm * 567))); grid.append(gc)
        for row in t.rows:
            tc = OxmlElement("w:tc"); tcpr = OxmlElement("w:tcPr")
            w = OxmlElement("w:tcW"); w.set(qn("w:w"), str(int(image_cm * 567))); w.set(qn("w:type"), "dxa")
            tcpr.append(w); tc.append(tcpr); tc.append(OxmlElement("w:p")); row._tr.append(tc)
        top = t.cell(0, 2).merge(t.cell(len(t.rows) - 1, 2))
        va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); top._tc.get_or_add_tcPr().append(va)
        para = top.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        para.add_run().add_picture(str(image), width=Cm(image_cm - 0.3))
        for row in t.rows:
            for cell in row.cells:
                for q in cell.paragraphs:
                    q.paragraph_format.keep_with_next = True
        if caption:
            self.caption("Figure", caption)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def image(self, path, width_cm, caption=None):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        para.paragraph_format.keep_with_next = True
        if caption:
            self.caption("Figure", caption)

    def images_row(self, paths_and_labels, width_cm, caption=None):
        """Screenshots side by side in a borderless table."""
        n = len(paths_and_labels)
        t = self.d.add_table(rows=2, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (path, label) in enumerate(paths_and_labels):
            c = t.rows[0].cells[i]; para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(path), width=Cm(width_cm))
            lp = t.rows[1].cells[i].paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = lp.add_run(label); r.font.size = Pt(8); r.italic = True
        if caption:
            self.caption("Figure", caption)
        else:
            self.d.add_paragraph()

    def save(self, path):
        self.d.save(str(path))
